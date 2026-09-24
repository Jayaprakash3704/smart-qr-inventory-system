import argparse
import os
from datetime import datetime

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


def configure_page(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.orientation = WD_ORIENT.PORTRAIT


def set_default_font(doc: Document, font_name: str = "Calibri", size: int = 11) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = font_name
    normal.font.size = Pt(size)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)


def add_cover_page(
    doc: Document,
    company_name: str,
    prepared_by: str,
    version: str,
    app_name: str,
    generated_on: str,
) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run = p.add_run(company_name)
    run.bold = True
    run.font.size = Pt(24)
    run.font.color.rgb = RGBColor(31, 41, 55)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run(f"User Manual for {app_name}")
    run.bold = True
    run.font.size = Pt(18)
    run.font.color.rgb = RGBColor(17, 24, 39)

    doc.add_paragraph()

    meta = doc.add_table(rows=5, cols=2)
    meta.style = "Table Grid"
    meta.cell(0, 0).text = "Document Title"
    meta.cell(0, 1).text = f"{app_name} - Detailed Operations Manual"
    meta.cell(1, 0).text = "Version"
    meta.cell(1, 1).text = version
    meta.cell(2, 0).text = "Prepared For"
    meta.cell(2, 1).text = company_name
    meta.cell(3, 0).text = "Prepared By"
    meta.cell(3, 1).text = prepared_by
    meta.cell(4, 0).text = "Date"
    meta.cell(4, 1).text = generated_on

    for row in meta.rows:
        for idx, cell in enumerate(row.cells):
            for par in cell.paragraphs:
                for r in par.runs:
                    if idx == 0:
                        r.bold = True

    doc.add_paragraph()
    note = doc.add_paragraph(
        "Confidential: This document provides detailed, step-by-step workflow instructions and operational guides for all functional modules of the system."
    )
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    note.runs[0].italic = True

    doc.add_page_break()


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    h = doc.add_heading(text, level=level)
    if h.runs:
        h.runs[0].font.color.rgb = RGBColor(17, 24, 39)


def add_bullet_points(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(item, style="List Bullet")
        p.paragraph_format.space_after = Pt(3)

def add_screenshot(doc: Document, img_path: str, title: str, purpose: str, step_instructions: list[str], level: int = 2):
    add_heading(doc, title, level=level)
    
    # Purpose
    p_purpose = doc.add_paragraph()
    r1 = p_purpose.add_run("Overview & Purpose: ")
    r1.bold = True
    p_purpose.add_run(purpose)
    
    # Steps
    if step_instructions:
        p_steps = doc.add_paragraph()
        r2 = p_steps.add_run("Step-by-Step Instructions:")
        r2.bold = True
        add_bullet_points(doc, step_instructions)
    
    if os.path.exists(img_path):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run()
        r.add_picture(img_path, width=Inches(6.5))
        p.paragraph_format.space_after = Pt(20)
    else:
        doc.add_paragraph(f"[Screenshot temporarily missing or unverified: {img_path}]", style="Intense Quote")
    
def build_manual(
    output: str,
    company_name: str,
    prepared_by: str,
    version: str,
    app_name: str,
    screenshots_dir: str
) -> None:
    doc = Document()
    configure_page(doc)
    set_default_font(doc)

    generated_on = datetime.now().strftime("%d %B %Y")

    add_cover_page(doc, company_name, prepared_by, version, app_name, generated_on)
    
    add_heading(doc, "1. Executive Summary", level=1)
    doc.add_paragraph(
        "Sudhan Textile ERP is a comprehensive and advanced management solution tailored strictly for the textile manufacturing industry, "
        "focusing profoundly on Sizing and Warping operations. The system helps manage the core foundation via Masters, controls transactions via Sizing ERP, "
        "and handles global logic through System Settings. This manual offers deep operational documentation on utilizing the full web-based software suite, covering all 27 distinct views."
    )

    doc.add_page_break()

    # Dashboard
    add_heading(doc, "2. Dashboard & System Analytics", level=1)
    add_screenshot(
        doc, os.path.join(screenshots_dir, "1.png"), 
        "2.1 Administrative Dashboard", 
        "The Dashboard serves as the central hub of operations, delivering immediate analytical insights on plant performance and pending tasks.",
        [
            "View aggregated metrics such as Total Active Sizing Sets, Today's Sizing Production, Total Yarn Stock, and Invoices Pending Payment.",
            "Use the visual graphs and efficiency trackers to monitor the health of current mechanical operations.",
            "Utilize the quick-access actions integrated on the right-hand panel to immediately jump into creating a New Set or receiving new Yarn."
        ]
    )
    doc.add_page_break()

    # Masters Module
    add_heading(doc, "3. System Masters Module", level=1)
    doc.add_paragraph("Master configurations are the foundational data points used to populate dropdowns, define parameters, and standardize metrics across operational forms.")
    
    add_screenshot(
        doc, os.path.join(screenshots_dir, "2.png"), 
        "3.1 Company Master", 
        "Central configuration defining your organizational identity and formal contact information.",
        [
            "Navigate to the 'Masters' section and select 'Company'.",
            "Fill in essential fields including the Official Company Name, Registration details (PAN/GSTIN), and precise communication addresses.",
            "Save the record. Note that details entered here directly reflect on generated Tax Invoices and Delivery Challans."
        ]
    )

    add_screenshot(
        doc, os.path.join(screenshots_dir, "3.png"), 
        "3.2 Parties Master", 
        "A singular repository managing all external entities the business interfaces with, including suppliers, clients, and partner mills.",
        [
            "Open the 'Parties Master' dashboard.",
            "Click on 'Add New Party' to register a new vendor or client.",
            "Specify their category (e.g., Weaving Mill) to appropriately restrict their assignment to job cards.",
            "Update contact properties and formal billing details like their specific tax boundaries."
        ]
    )

    add_screenshot(
        doc, os.path.join(screenshots_dir, "4.png"), 
        "3.3 Yarn Count Master", 
        "A critical matrix establishing the varied specifications of raw yarn brought into the factory.",
        [
            "Navigate to 'Yarn Count Master'.",
            "Introduce standard designations (e.g., '20s 2/100', '30s 2/80').",
            "Attach intrinsic properties like ply thickness. These pre-configured counts will speed up the process during Yarn inward documentation."
        ]
    )

    add_screenshot(
        doc, os.path.join(screenshots_dir, "5.png"), 
        "3.4 Loom Type Master", 
        "Mapping of available manufacturing hardware capabilities ensuring job cards align with physical capacities.",
        [
            "Access the 'Loom Type Master' menu.",
            "Register hardware models such as 'Sulzer' or 'Air Jet'.",
            "Define the maximum physical width in inches to enforce capacity rules during production allocations."
        ]
    )

    add_screenshot(
        doc, os.path.join(screenshots_dir, "6.png"), 
        "3.5 Beam Master", 
        "An active registry tracking the lifecycle and mechanical properties of physical beams.",
        [
            "Open 'Beam Master' to see a breakdown of 'Sizing Beams' and 'Warping Beams'.",
            "Register individual beams with unique Beam Numbers.",
            "Configure their tare weights (kg) and maximum ends. This data is critical for net yarn weight calculations."
        ]
    )

    add_screenshot(
        doc, os.path.join(screenshots_dir, "7.png"), 
        "3.6 Vehicle Master", 
        "Logistical management of the active transport fleet bringing goods in and out.",
        [
            "Add logistics vectors by entering the Vehicle Registration Numbers.",
            "Maintain details regarding primary drivers and their contact coordinates.",
            "These vehicles are subsequently linked to Delivery Challans and Inbound Receipts."
        ]
    )

    add_screenshot(
        doc, os.path.join(screenshots_dir, "8.png"), 
        "3.7 Financial Year Master", 
        "Time-boxing parameters configuring sequential numbering rules and financial closure calculations.",
        [
            "Head to 'Financial Year Master' configuration.",
            "Establish absolute Start Dates and End Dates for the fiscal cycle.",
            "Activating a financial year directly impacts the Document Series suffixes applied system-wide."
        ]
    )

    add_screenshot(
        doc, os.path.join(screenshots_dir, "9.png"), 
        "3.8 Document Series Master", 
        "Regulation of automatic identification prefixing and invoice enumerations to maintain audit integrity.",
        [
            "Under Document Series, configure the string Prefixes for different forms (e.g., 'INV' for Invoices).",
            "Link them to the Active Financial Year to generate standardized formats like 'INV0001/2025-26'."
        ]
    )

    doc.add_page_break()

    # Sizing ERP Module
    add_heading(doc, "4. Sizing ERP Operations", level=1)
    doc.add_paragraph("This functional core tracks every phase of the physical textile operation—from raw yarn intake, machine processing, right down to delivery and invoicing.")

    add_screenshot(
        doc, os.path.join(screenshots_dir, "10.png"), 
        "4.1 Yarn Receipt (Inward Workflow)", 
        "Formal tracking protocol for all bound raw materials entering the warehouse.",
        [
            "Navigate to Sizing ERP > Yarn Receipt.",
            "Click on '+ New Receipt' and fill in the DC number assigned by the transport provider.",
            "Select the delivering Party and assign the internal vehicle associated with the drop.",
            "Add line specific items: defining exact yarn count lots, cone frequencies, and weight measurements.",
            "Submit the form moving it to 'Pending Approval' for inventory update authorization."
        ]
    )
    
    add_screenshot(
        doc, os.path.join(screenshots_dir, "11.png"), 
        "4.2 Baby Cone / Winding Operations", 
        "Managing intermediary conversion procedures dividing the major cones into manufacturing appropriate sizing sets.",
        [
            "Access the Winding forms.",
            "Record inbound bulk cones and lot classifications.",
            "Measure and enter the calculated quantity, output weights of the newly wound 'Baby Cones'.",
            "This inherently logs percentage waste parameters during process transition."
        ]
    )

    add_screenshot(
        doc, os.path.join(screenshots_dir, "12.png"), 
        "4.3 Warping Job Card Initiation", 
        "Orchestrating primary weaving preparation steps by designating batches to warping equipment.",
        [
            "Go to 'Warping Job Card' under Sizing ERP.",
            "Initialize a new job card linked to a specific Party requesting the warp.",
            "Assign detailed Yarn lots to corresponding machines.",
            "Set objective constraints like Length target parameters required per job configuration."
        ]
    )

    add_screenshot(
        doc, os.path.join(screenshots_dir, "13.png"), 
        "4.4 Sizing Job Card (Set Report)", 
        "High-level administration for end-to-end sizing production schedules.",
        [
            "View all pending, drafted, or authorized sizing sequences from the Job Card dashboard.",
            "Click 'New Set' to define a consolidated sizing operation encompassing machine type (Loom), targeted ends, and meter goals.",
            "Monitor status indicators showing physical stages and pending authorizations before the set goes active."
        ]
    )

    add_screenshot(
        doc, os.path.join(screenshots_dir, "14.png"), 
        "4.5 Beam Management Console", 
        "Dynamic tracking of the entire available inventory of functional factory beams.",
        [
            "Open 'Beam Management'.",
            "Observe the status matrix distinguishing 'Available' beams ready for assignment vs. 'In Process' beams mounted within an active Sizing Card.",
            "Review tare metrics and use tracking tags connecting beams back to live processes."
        ]
    )

    add_screenshot(
        doc, os.path.join(screenshots_dir, "15.png"), 
        "4.6 Yarn Stock Ledger", 
        "A highly-reliable, read-only auditing matrix tabulating exact global and segmented inventory amounts.",
        [
            "Access the 'Yarn Stock Ledger'.",
            "Use filters to condense views by distinct Yarn Count, Source Party, or particular shipment Lot Numbers.",
            "Analyze metrics such as 'Total Inward (kg)' versus 'Total Outward (kg)' calculating definitive real-time 'Balance Stock'.",
            "Export this ledger externally (PDF/CSV) for weekly physical audit operations."
        ]
    )

    add_screenshot(
        doc, os.path.join(screenshots_dir, "16.png"), 
        "4.7 Yarn Return DC Creation", 
        "Processing protocol for returning flawed or surplus materials to original vendor mills.",
        [
            "Navigate to 'Yarn Return'.",
            "Create a new document to explicitly log job-work returns preventing balance discrepancies.",
            "Specify exact return weights referencing the initial Lot Numbers."
        ]
    )

    add_screenshot(
        doc, os.path.join(screenshots_dir, "17.png"), 
        "4.8 Yarn Delivery DC Compilation", 
        "Final dispatch procedural records clearing operations upon project completion.",
        [
            "Move to the 'Yarn Delivery' component.",
            "Draft a new Outbound Delivery pointing to finalized Sizing Sets or Beams.",
            "Verify total gross vs net weight (taring) before assigning to a delivery vehicle, generating a formal outbound challan document."
        ]
    )

    add_screenshot(
        doc, os.path.join(screenshots_dir, "18.png"), 
        "4.9 GST Tax Invoice Finalization", 
        "Conversion of operational completed data into binding financial accounts.",
        [
            "Access 'GST Tax Invoice'.",
            "Click '+ New Invoice'. You can typically import verified data spanning from an authorized Delivery DC.",
            "Let the system automatically calculate defined Service amounts and HSN/SAC parameters.",
            "Verify the automated split of Tax components (CSGT, SGST, IGST) determined by inter-state rules.",
            "Save and export the invoice for the finance team/party delivery."
        ]
    )

    doc.add_page_break()

    # Settings & Security Module
    add_heading(doc, "5. Settings, Security & Governance", level=1)
    doc.add_paragraph("This sector maintains the granular configuration governing security, data longevity, logic compliance, and operational rights matrices limits.")

    add_screenshot(
        doc, os.path.join(screenshots_dir, "19.png"), 
        "5.1 Profile Settings Initialization", 
        "Allows individual operators to modify their localized contact profiles.",
        [
            "Access the 'Profile Settings' page through the settings menu.",
            "Modify personal names and review account readiness.",
            "Review localized metrics concerning 'Last Login' timestamping and current system role definition."
        ]
    )

    add_screenshot(
        doc, os.path.join(screenshots_dir, "20.png"), 
        "5.2 User Management", 
        "SuperAdmin control module orchestrating all human interaction within the application.",
        [
            "Open 'User Management'.",
            "Click '+ Add User' to assign credentials for new personnel.",
            "Provide explicit Roles (Manager, Operator) linking their identity to respective departments.",
            "Administratively override, lock, or disable users suspected of invalid activities."
        ]
    )

    add_screenshot(
        doc, os.path.join(screenshots_dir, "21.png"), 
        "5.3 Role & Permission Configuration", 
        "Implementation of strict Role-Based Access Control logic.",
        [
            "Navigate to the 'Role Permissions' pane.",
            "Select specific predefined roles (e.g., Operator layer versus Managerial layer).",
            "Activate or retract detailed module-specific checkboxes defining Read, Write, and Authorization capabilities per form type."
        ]
    )

    add_screenshot(
        doc, os.path.join(screenshots_dir, "22.png"), 
        "5.4 Approval Matrix Structuring", 
        "Configuration defining secondary verification thresholds on critical documents.",
        [
            "Go to 'Approval Matrix'.",
            "Select distinct workflows (e.g., 'Yarn Receipt' vs 'Sizing Job Card').",
            "Configure mandatory authorization limits. Enable configurations dictating that Job Cards exceeding predetermined weight or value thresholds demand explicit managerial digital signatures before processing."
        ]
    )

    add_screenshot(
        doc, os.path.join(screenshots_dir, "23.png"), 
        "5.5 System Configuration Constants", 
        "Root level deployment configuration defining baseline logical values.",
        [
            "Check 'System Settings'.",
            "Implement standardized values ensuring consistency, such as global 'Company Currency' to 'INR'.",
            "Fix the 'Time Zone' variable guaranteeing accurate timestamp correlations required for precise database transactions."
        ]
    )

    add_screenshot(
        doc, os.path.join(screenshots_dir, "24.png"), 
        "5.6 Security Policies Enactment", 
        "Risk mitigation configurations to throttle external attacks.",
        [
            "Open 'Security Policies'.",
            "Establish lockout thresholds mitigating brute force behavior by limiting 'Max Failed Login Attempts' (e.g., set to 5).",
            "Designate the strict automated 'Lockout Duration' forcing suspended timeouts."
        ]
    )

    add_screenshot(
        doc, os.path.join(screenshots_dir, "26.png"), 
        "5.7 Audit Logs Compliance Tracker", 
        "Immutable chronological surveillance of all recorded systemic transactions.",
        [
            "View 'Audit Logs' to analyze historical operational flow.",
            "Filter through chronological stamps highlighting specific 'Actions' (Create, Update, Delete).",
            "Cross-references data points tracking modifications tied back directly to specific IP addresses and internal Users preventing operational denials."
        ]
    )

    add_screenshot(
        doc, os.path.join(screenshots_dir, "27.png"), 
        "5.8 Backup & Data Safety Scheduling", 
        "Routine prevention ensuring operational integrity via database snapshots.",
        [
            "Open 'Backup & Safety'.",
            "Configure interval timings scheduling automated 'Full Database Backups'.",
            "Define retention policies dictating how many days historically the local files are maintained.",
            "Execute out-of-cycle saves utilizing the manual 'Run Backup Now' function."
        ]
    )

    add_screenshot(
        doc, os.path.join(screenshots_dir, "28.png"), 
        "5.9 Notification & Alert Configuration", 
        "Systemic broadcasting ensuring important process failures alert necessary personnel.",
        [
            "Navigate to 'Notifications'.",
            "Turn on explicit systemic events (e.g., flagging 'Backup Failure Alerts').",
            "Select which functional Roles (e.g. Administrators only) are pushed the alerts.",
            "Initiate mock tests triggering the alert network mapping verifications."
        ]
    )

    doc.save(output)

if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", default="Sudhan_Textile_ERP_User_Manual.docx")
    parser.add_argument("--company", default="Sudhan Textile Mills")
    parser.add_argument("--prepared-by", default="Implementation Team")
    parser.add_argument("--version", default="1.0")
    parser.add_argument("--app-name", default="Sudhan Textile ERP")
    parser.add_argument("--screenshots", default="Screenshots", help="Directory containing screenshots")
    
    args = parser.parse_args()
    build_manual(
        output=args.output,
        company_name=args.company,
        prepared_by=args.prepared_by,
        version=args.version,
        app_name=args.app_name,
        screenshots_dir=args.screenshots
    )
    print(f"ERP user manual created successfully: {args.output}")
